"""BOLO 83 — build one tailored résumé (.docx) from a JSON spec. ATS-safe: one column, no tables, no images.

Usage:  .venv/Scripts/python _tools/bolo83/resume.py <spec.json> <out.docx>
Spec:   {name, contact, summary, skills: [..], experience: [{title, org, where, dates, bullets: [..]}],
         education: [..], certs: [..]}
Specs and output live in _PRIVATE/bolo83-applications/ (the repo is public).
"""
import json, sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

spec = json.load(open(sys.argv[1], encoding="utf-8"))
doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(0.6)
    s.left_margin = s.right_margin = Inches(0.7)
base = doc.styles["Normal"]; base.font.name = "Calibri"; base.font.size = Pt(10.5)
base.paragraph_format.space_after = Pt(2)

def para(text="", bold=False, size=None, align=None, after=2, italic=False):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    if align: p.alignment = align
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    return p

def head(text):
    p = para(text.upper(), bold=True, size=11, after=3); p.paragraph_format.space_before = Pt(8)

para(spec["name"], bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para(spec["contact"], align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
head("Summary"); para(spec["summary"])
head("Core skills"); para(" · ".join(spec["skills"]))
head("Experience")
for j in spec["experience"]:
    p = para(after=0); p.paragraph_format.space_before = Pt(5)
    r = p.add_run(f'{j["title"]} — {j["org"]}'); r.bold = True
    p.add_run(f'  |  {j["where"]}  |  {j["dates"]}')
    for b in j["bullets"]:
        bp = doc.add_paragraph(b, style="List Bullet"); bp.paragraph_format.space_after = Pt(1)
head("Education")
for e in spec["education"]: para(e)
head("Certifications")
for c in spec["certs"]: para(c)
doc.save(sys.argv[2]); print("wrote", sys.argv[2])
